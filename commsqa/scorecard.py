"""Excel QA scorecard with native Excel charts (openpyxl)."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import qa
from .rules import RULES_BY_CODE

HEAD_FILL = PatternFill("solid", fgColor="DDE3EA")
BOLD = Font(bold=True)


def _table(ws, top_row: int, headers: list, rows: list, formats: dict | None = None) -> int:
    for j, h in enumerate(headers, 1):
        c = ws.cell(top_row, j, h)
        c.font, c.fill = BOLD, HEAD_FILL
    for i, row in enumerate(rows, top_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            if formats and headers[j - 1] in formats:
                c.number_format = formats[headers[j - 1]]
    return top_row + len(rows)


def _autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(9, width + 2), 70)


def build(conn, path) -> None:
    run_id = qa.latest_run_id(conn)
    k = qa.kpis(conn, run_id)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Client correspondence QA scorecard"
    ws["A1"].font = Font(bold=True, size=13)
    _table(ws, 3, ["Measure", "Value"], [
        ["Documents checked (100%)", k["documents"]], ["Defective documents", k["defective_documents"]],
        ["Accuracy", k["accuracy"]], ["Defects", k["defects"]], ["Critical defects", k["critical"]],
        ["Open for review", k["open"]]])
    ws["B6"].number_format = "0.00%"
    end = _table(ws, 11, ["Root cause", "Defects", "Documents"],
                 [[r["label"], r["n"], r["docs"]] for r in qa.breakdown(conn, run_id, "root_cause")])
    _table(ws, end + 3, ["Document type", "Documents", "Defective", "Accuracy"],
           [[r["doc_type"], r["documents"], r["defective"], r["accuracy"]] for r in qa.by_doc_type(conn, run_id)],
           {"Accuracy": "0.00%"})
    _autosize(ws)

    ws = wb.create_sheet("Weekly trend")
    weeks = qa.weekly(conn, run_id)
    last = _table(ws, 1, ["Week", "Documents", "Defective", "Defects", "Critical", "Accuracy"],
                  [[w["week"], w["documents"], w["defective"], w["defects"], w["critical"], w["accuracy"]] for w in weeks],
                  {"Accuracy": "0.00%"})
    bars = BarChart()
    bars.title, bars.y_axis.title = "Defective documents per week", "Documents"
    bars.add_data(Reference(ws, min_col=3, min_row=1, max_row=last), titles_from_data=True)
    bars.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    bars.height, bars.width = 7.5, 16
    ws.add_chart(bars, "H2")
    line = LineChart()
    line.title, line.y_axis.title = "Accuracy", "Accuracy"
    line.y_axis.number_format = "0.0%"
    line.add_data(Reference(ws, min_col=6, min_row=1, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    line.height, line.width = 7.5, 16
    ws.add_chart(line, "H18")
    _autosize(ws)

    ws = wb.create_sheet("Pareto")
    rows = qa.pareto(conn, run_id)
    last = _table(ws, 1, ["Rule", "Title", "Severity", "Root cause", "Defects", "Share", "Cumulative"],
                  [[r["rule"], r["title"], r["severity"], r["root_cause"], r["n"], r["share"], r["cumulative"]] for r in rows],
                  {"Share": "0.0%", "Cumulative": "0.0%"})
    chart = BarChart()
    chart.title = "Defects by rule"
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=last))
    chart.height, chart.width = 9, 20
    ws.add_chart(chart, "I2")
    _autosize(ws)

    ws = wb.create_sheet("Defect log")
    log = conn.execute("""SELECT x.defect_id, x.doc_id, d.doc_type, d.week, d.issued_on, x.rule, x.severity,
                                 x.root_cause, x.field, x.expected, x.found, x.status, x.reviewer, x.comment
                          FROM defects x JOIN documents d ON d.doc_id = x.doc_id
                          WHERE x.run_id = ? ORDER BY x.defect_id""", (run_id,)).fetchall()
    headers = ["Defect", "Document", "Type", "Week", "Issued", "Rule", "Severity", "Root cause", "Field",
               "Expected", "Found", "Status", "Reviewer", "Comment"]
    last = _table(ws, 1, headers, [list(r) for r in log])
    for r in range(2, last + 1):
        ws.cell(r, 6).value = f"{ws.cell(r, 6).value} {RULES_BY_CODE[ws.cell(r, 6).value].title}"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = "A2"
    _autosize(ws)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.save(path)
